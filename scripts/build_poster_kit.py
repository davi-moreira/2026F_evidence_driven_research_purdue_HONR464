#!/usr/bin/env python3
"""build_poster_kit.py — the Expo poster template and the poster rubric (D87).

Two artifacts the syllabus promised ("a poster template and assessment rubric
will be shared", D53) and that the D50 conference block left open:

  project/poster/honr464_poster_template.pptx
      The research poster template, adapted from QM 47400's business-analytics
      template (qm474_poster_template.pptx) to the EDR|AI way of doing research.
      48 x 36 in, the Purdue Fall Undergraduate Research Expo's standard 4' x 3'
      board. Same visual family as QM 47400's (grey outer columns, black bars
      with Purdue-gold headings, Arial), because both courses print in one run;
      the panels follow the book instead of CRISP-DM. A headline banner carries
      the finding and its qualifier in one eye-span (ch. 32); three unequal
      columns give the main result the centre: the question and what it adds
      (ch. 7-9), the compass declaration and its warrants (ch. 6-7, 10, 12,
      14-19) and the data card (ch. 20-21) on the left; the main result and the
      checks (ch. 11, 22-25, 32, 36-37) in the centre; what the evidence
      supports and where it stops (ch. 28-29, 35), then sources, on-poster AI
      disclosure and credits (ch. 9, 30) on the right. Slide 2 is a HIDDEN
      guide page (print spec, rules, panel -> rubric -> chapter map). The
      layout was re-hierarchized after the Codex review of 2026-10-21.

  project/poster/Evidence_Driven_Research_Poster_Rubric.{docx,pdf,xlsx}
      The final-poster rubric, adapted from QM 47400's poster-only rubric
      (Predictive_Analytics_Poster_Rubric) to research: eight criteria, one per
      panel group, four bands, 100 points, with the hard caps. It is THE
      Final Project poster-quality rubric: the same table is injected into the
      M13 brief between the `poster-rubric` markers, so the brief, the PDF the
      students hold and the grading sheet can never disagree.

The rubric text lives once, in RUBRIC / HARD_CAPS / CAP_NOTE below. The
template's example is EDR|AI ch. 32's simulated checkout comparison (SEED =
464), redrawn at print size with a second panel for the difference the
headline is about (Welch 95% interval, weeks treated as independent).

    .venv/bin/python scripts/build_poster_kit.py              # everything
    .venv/bin/python scripts/build_poster_kit.py --rubric     # rubric + M13 block only
    .venv/bin/python scripts/build_poster_kit.py --check      # exit 1 if M13's block is stale

The template build reads QM 47400's template for its Purdue logo (from the
sibling course folder, or --qm474 PATH). The rubric build needs nothing outside
this repository. PDFs are produced with LibreOffice (`soffice`).
"""
from __future__ import annotations

import argparse
import copy
import io
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "project" / "poster"
TEMPLATE_OUT = OUT_DIR / "honr464_poster_template.pptx"
RUBRIC_STEM = "Evidence_Driven_Research_Poster_Rubric"
M13 = ROOT / "_research_project" / "2026Fall" / "milestone_13_final_poster_lock.md"
QM474_DEFAULT = (
    ROOT.parents[1]
    / "predictive_analytics"
    / "2026F_predictive_analytics_QM474"
    / "_final_project"
    / "2026F"
    / "template"
    / "qm474_poster_template.pptx"
)
BEGIN = "<!-- poster-rubric:begin -->"
END = "<!-- poster-rubric:end -->"
SEED = 464

# ---------------------------------------------------------------------------
# The rubric (single source)
# ---------------------------------------------------------------------------
BANDS = ("Exemplary", "Proficient", "Developing", "Beginning")

RUBRIC = [
    {
        "name": "Question, context & contribution",
        "points": 10,
        "panels": "Title · The question and what it adds",
        "virtue": "Compass & pathway alignment",
        "chapters": "Ch. 1, 7–9",
        "bands": [
            ("9–10", "One specific research question, and why it matters; the main conclusion of the literature synthesized from at least two sources you read; what your project adds to that conclusion made explicit; the hypothesis you recorded before the analysis ran, with any later analysis labeled exploratory."),
            ("7–8", "The question is clear, but the context or the contribution stays general: one source, or \"little is known\" without saying what is known."),
            ("4–6", "The question is broad or ambiguous; sources are listed without saying what they establish; no contribution is stated."),
            ("0–3", "A topic without a usable research question."),
        ],
    },
    {
        "name": "Target, design & warranted claims",
        "points": 15,
        "panels": "Design: what your claim needs",
        "virtue": "Compass & pathway alignment",
        "chapters": "Ch. 6–7, 10, 12, 14–19",
        "bands": [
            ("14–15", "The question's kind (descriptive or causal) and its reach (a named population beyond the data, or the unseen cases a prediction serves) are declared, with prediction named as an objective when it is one; the target quantity is exact; the board states what supports each crossing (a sampling argument for a population, an assignment or identification argument with its assumptions for a cause, a held-out test against a baseline for prediction) and what the evidence cannot yet support, including a causal target that remains unidentified."),
            ("11–13", "Design and target are clear; one warrant or assumption is underexplained."),
            ("6–10", "The link between target, design, and reach is ambiguous, or the kind is mislabeled (for example, a causal question silently relabeled descriptive)."),
            ("0–5", "No defensible account of what the design can establish, or a design presented as able to answer a question it cannot."),
        ],
    },
    {
        "name": "Data, provenance & measurement",
        "points": 12,
        "panels": "Data and measurement",
        "virtue": "Evidence integrity & provenance",
        "chapters": "Ch. 13, 20–21",
        "bands": [
            ("11–12", "Source and route can be identified (archive, URL or DOI, retrieval date, or collection protocol); observations and independent units are distinguished, with any grouping named; the time window is given; each key concept is tied to its measure with the main validity concern; missing data and their handling are reported; the permission or IRB determination is stated."),
            ("9–10", "The evidence is identifiable; one consequential detail is thin."),
            ("5–8", "Important provenance or measurement information is missing."),
            ("0–4", "The evidence cannot be identified, or a material permission problem is unresolved."),
        ],
    },
    {
        "name": "Analysis, checks & verification",
        "points": 15,
        "panels": "Checks and verification · verification status line",
        "virtue": "Verification & robustness",
        "chapters": "Ch. 11, 22–25, 36–37",
        "bands": [
            ("14–15", "The analysis fits the design; what was declared in advance is kept apart from later changes, and each change is explained; the uncertainty method fits the data, computed at the level of the independent unit with any dependence named; at least one consequential robustness, sensitivity, or diagnostic check is reported with what it changed; the verification status is accurate and honestly labeled (an external cold run, or \"solo proxy; external cold run pending\")."),
            ("11–13", "Analysis and uncertainty are appropriate; the check or the verification account is thin."),
            ("6–10", "A weak link between method and design, an uninformative check, or a verification claim the record does not support."),
            ("0–5", "A material analytical error, or a displayed result with no trace or that fails to reproduce."),
        ],
    },
    {
        "name": "Results & figure honesty",
        "points": 15,
        "panels": "Main result",
        "virtue": "Verification & robustness",
        "chapters": "Ch. 32",
        "bands": [
            ("14–15", "The main evidence answers the question and is easy to inspect; key estimates and comparisons are labeled directly, with units and the n; the uncertainty the design warrants is shown (an interval with its type and level, or a precise caveat saying why none is warranted); scales are honest (bars start at zero, other plots carry clearly labeled scales); the caption names the notebook cell or output behind the figure."),
            ("11–13", "The result is understandable and honest; one labeling or caption omission."),
            ("6–10", "The reader has to reconstruct a key comparison, a denominator, or the uncertainty."),
            ("0–5", "The visual encoding distorts the result (a truncated bar axis turning a small gap into a cliff), or no inspectable result is presented."),
        ],
    },
    {
        "name": "Interpretation, uncertainty & limitations",
        "points": 15,
        "panels": "Headline · What this means, and where it stops",
        "virtue": "Uncertainty & limitations",
        "chapters": "Ch. 28–29, 35",
        "bands": [
            ("14–15", "Headline and conclusion claim no more than the evidence: causal wording only with a defended identification argument, population wording only with a reach argument; the qualifier that blocks the likely wrong inference sits in the same eye-span as the claim; the consequential limitation and a feasible next step are explicit (uncertainty, limitation, next step); the preliminary status is stated."),
            ("11–13", "The interpretation is sound; the qualifier's placement or the next step is underdeveloped."),
            ("6–10", "Generic limitations (\"results may vary\"), uncertainty off the read path, or wording that drifts past the evidence."),
            ("0–5", "The central claim exceeds its warrant, or the result is presented as settled certainty."),
        ],
    },
    {
        "name": "Sources, AI disclosure & credit",
        "points": 8,
        "panels": "Sources, AI use, and credits",
        "virtue": "Evidence integrity & provenance",
        "chapters": "Ch. 9, 30",
        "bands": [
            ("8", "Citations and references correspond and survive checking, and each reference can be found (author, year, title, venue, DOI or URL); the AI disclosure is on the poster itself and names each tool, its purpose, what you changed or kept from its output, and how you checked it (or \"No AI used\", when true); help, data providers, and funding are credited."),
            ("6–7", "Complete, with one reference or one disclosure line ambiguous."),
            ("3–5", "References or disclosure incomplete enough to impede checking."),
            ("0–2", "A fabricated or materially misattributed source, or a confirmed AI contribution missing from the disclosure."),
        ],
    },
    {
        "name": "Visual communication & accessibility",
        "points": 10,
        "panels": "The whole board",
        "virtue": "Craft & communication",
        "chapters": "Ch. 32–33",
        "bands": [
            ("9–10", "A clear reading order: the finding is found in seconds, and the question, result, and boundary in a ninety-second read; built on the course template at 48 × 36 in, with the guide's minimum type sizes; strong contrast; every distinction carried in two channels (a label or pattern, never colour alone); each figure carries alt text in the file; no template guidance or EXAMPLE material left."),
            ("7–8", "Readable and navigable, with minor craft lapses or one colour-only distinction."),
            ("4–6", "Dense, small, or poorly ordered content that needs the presenter; template text left on the board."),
            ("0–3", "The main message is illegible at print size, or meaning is carried by colour alone throughout."),
        ],
    },
]

HARD_CAPS = [
    ("A fabricated source, a source that cannot be identified after reasonable checking, or a materially "
     "false attribution to a source", "Sources, AI disclosure & credit"),
    ("A displayed result your project computed, with no reproducible trace to a notebook cell or output",
     "Analysis, checks & verification"),
    ("A displayed result or bound that materially fails to reproduce from the submitted package "
     "(documented rounding is not a failure)", "Analysis, checks & verification"),
    ("An author rerun presented as independent reproduction, or a solo proxy without its label",
     "Analysis, checks & verification"),
    ("A central claim that materially exceeds its identification or reach warrant",
     "Interpretation, uncertainty & limitations"),
    ("A confirmed AI contribution missing from the on-poster disclosure", "Sources, AI disclosure & credit"),
]

CAP_NOTE = (
    "Apply each cap once, to its named criterion: the score falls to the top of the Beginning band, "
    "or stays at the lower score already earned. A temporary link failure is not a fabricated source, "
    "and missing ledger rows do not by themselves prove undisclosed AI use. Record the evidence behind "
    "any cap in the feedback."
)

assert sum(c["points"] for c in RUBRIC) == 100, "rubric must total 100"


def rubric_markdown() -> str:
    """The table + caps block injected into the M13 brief."""
    lines = [
        BEGIN,
        "<!-- Generated by scripts/build_poster_kit.py (D87) from its RUBRIC and",
        "     HARD_CAPS. Edit the script and rerun it; edits here are overwritten. -->",
        "",
        "Score each criterion in one band. The rubric scores what your poster",
        "communicates, panel by panel; the grader checks that it is true (sources,",
        "computed results, reproduction, AI disclosure) against your M9 note, your",
        "package, your run records, and your ledger. Those checks decide whether the",
        "board can be trusted; they add no package-craft points. The same rubric is",
        f"handed out as `project/poster/{RUBRIC_STEM}.pdf` (with the grading sheet",
        f"`{RUBRIC_STEM}.xlsx`) next to the poster template",
        "`project/poster/honr464_poster_template.pptx`; each criterion names the",
        "template panels it reads.",
        "",
        "| # | Criterion | " + " | ".join(BANDS) + " |",
        "|---|---|---|---|---|---|",
    ]
    for i, c in enumerate(RUBRIC, 1):
        cells = [
            f"{text} ({rng})" for rng, text in c["bands"]
        ]
        crit = f"**{c['name']}** ({c['points']})<br>*{c['panels']}*"
        lines.append(f"| {i} | {crit} | " + " | ".join(cells) + " |")
    lines += [
        "",
        "**Hard caps (a single failure caps the criterion at Beginning, regardless of the rest):**",
        "",
    ]
    for what, row in HARD_CAPS:
        lines.append(f"- {what}: caps *{row}* at Beginning.")
    lines += ["", CAP_NOTE, "", END]
    return "\n".join(lines)


def inject_m13(check: bool = False) -> bool:
    text = M13.read_text(encoding="utf-8")
    if BEGIN not in text or END not in text:
        sys.exit(f"{M13.name}: the {BEGIN} / {END} markers are missing")
    head, rest = text.split(BEGIN, 1)
    _, tail = rest.split(END, 1)
    new = head + rubric_markdown() + tail
    if check:
        return new == text
    if new != text:
        M13.write_text(new, encoding="utf-8")
    return True


# ---------------------------------------------------------------------------
# Rubric documents: .docx (and its PDF) and the .xlsx grading sheet
# ---------------------------------------------------------------------------
def build_rubric_docx(path: Path) -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(14), Inches(8.5)  # US legal, landscape
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(0.5))
    sec.top_margin = sec.bottom_margin = Inches(0.45)

    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    style.font.size = Pt(8.5)
    style.paragraph_format.space_after = Pt(2)

    def para(text, size=8.5, bold=False, italic=False, color=None, after=2):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.size, r.bold, r.italic = Pt(size), bold, italic
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        p.paragraph_format.space_after = Pt(after)
        return p

    para("Evidence-Driven Research Poster Rubric (Final Poster Assessment)", 14, bold=True, after=1)
    para(
        "HONR 46400 · Purdue Fall Undergraduate Research Expo · grades the locked poster "
        "(M13) as the Final Project poster-quality score: 70% of the Poster Presentation "
        "component, with the M15 live presentation supplying the other 30%. Built for the "
        "course template honr464_poster_template.pptx (48 × 36 in). Score each criterion in "
        "one band; the hard caps below override the bands. The rubric scores what the poster "
        "communicates; check that it is true against the M9 note, the package, the run records "
        "and the ledger.",
        8.5, italic=True, color="404040", after=4,
    )

    headers = ["#", "Criterion · panels graded", "Pts"] + list(BANDS) + ["Score"]
    widths = [0.3, 1.9, 0.4, 3.3, 2.2, 2.2, 2.0, 0.6]
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))

    def shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hexcolor)
        tcPr.append(shd)

    def fill(cell, runs, size=8):
        cell.text = ""
        p = cell.paragraphs[0]
        for text, bold, italic in runs:
            r = p.add_run(text)
            r.font.size, r.bold, r.italic = Pt(size), bold, italic
        p.paragraph_format.space_after = Pt(0)

    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]
        fill(c, [(h, True, False)], 8.5)
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string("CFB991")
        shade(c, "000000")
    # repeat the header row on every page
    trPr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    trPr.append(hdr)

    for i, crit in enumerate(RUBRIC, 1):
        cells = t.add_row().cells
        fill(cells[0], [(str(i), True, False)])
        fill(cells[1], [(crit["name"], True, False), ("\n" + crit["panels"], False, True),
                        (f"\n{crit['virtue']} · {crit['chapters']}", False, False)])
        fill(cells[2], [(str(crit["points"]), True, False)])
        for k, (rng, text) in enumerate(crit["bands"]):
            fill(cells[3 + k], [(f"{rng} pts: ", True, False), (text, False, False)])
        fill(cells[7], [("", False, False)])
        shade(cells[1], "EDEDED")
    total = t.add_row().cells
    fill(total[1], [("Total", True, False)])
    fill(total[2], [("100", True, False)])
    for row in t.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = Inches(w)
        # keep each criterion on one page
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        cant.set(qn("w:val"), "true")
        trPr.append(cant)

    para("", after=2)
    para("Hard caps: a single failure caps the criterion at Beginning, regardless of the rest", 9, bold=True, after=1)
    for what, row in HARD_CAPS:
        p = para("", after=0)
        p.style = doc.styles["List Bullet"]
        r = p.add_run(f"{what} → {row}.")
        r.font.size = Pt(8.5)
    para(CAP_NOTE, 8.5, after=2)
    para(
        "Separate from this score: the poster does not lock without the two blocking gate "
        "records (author self-reproduction log and release preflight), and a missing AI "
        "Research Ledger is handled in the M13 process score. For an approved group the "
        "poster-quality score is common; every member presents and receives their own live "
        "score. Source of truth: milestone_13_final_poster_lock.md (generated block).",
        8, italic=True, color="404040",
    )
    doc.core_properties.title = "Evidence-Driven Research Poster Rubric"
    doc.core_properties.author = "Davi Moreira"
    doc.save(path)


def build_rubric_xlsx(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Rubric"
    head = ["#", "Criterion", "Panels graded", "Points"] + [f"{b}" for b in BANDS] + ["Score", "Comment"]
    ws.append(head)
    thin = Side(style="thin", color="999999")
    for i, c in enumerate(RUBRIC, 1):
        ws.append([i, c["name"], c["panels"], c["points"]]
                  + [f"{rng} pts: {text}" for rng, text in c["bands"]] + [None, None])
    n = len(RUBRIC)
    ws.append([None, "Total", None, f"=SUM(D2:D{n + 1})"] + [None] * 4 + [f"=SUM(I2:I{n + 1})", None])
    ws.append([])
    ws.append([None, "Hard caps (cap the named criterion at Beginning):"])
    for what, row in HARD_CAPS:
        ws.append([None, what, row])
    ws.append([None, CAP_NOTE])
    widths = {"A": 4, "B": 30, "C": 30, "D": 8, "E": 55, "F": 45, "G": 45, "H": 45, "I": 9, "J": 40}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=1, max_row=n + 2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="CFB991")
        cell.fill = PatternFill("solid", fgColor="000000")
    for cell in ws[n + 2]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "B2"
    wb.save(path)


def to_pdf(src: Path, outdir: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    with tempfile.TemporaryDirectory() as profile:
        subprocess.run(
            [soffice, f"-env:UserInstallation=file://{profile}", "--headless",
             "--convert-to", "pdf", "--outdir", str(outdir), str(src)],
            check=True, capture_output=True,
        )
    return outdir / (src.stem + ".pdf")


# ---------------------------------------------------------------------------
# The poster template
# ---------------------------------------------------------------------------
INK, GOLD, GREY, WHITE = "000000", "CFB991", "E3E3E3", "FFFFFF"
CRIMSON, MUTED = "8A1C2B", "404040"
FONT = "Arial"
W_IN, H_IN = 48, 36
MARGIN, GUTTER = 0.75, 0.75
# Three unequal columns: the result takes the centre and most of the room (Codex review, D87).
COL_W = [13.0, 19.0, 13.0]
COL_X = [MARGIN, MARGIN + COL_W[0] + GUTTER, MARGIN + COL_W[0] + COL_W[1] + 2 * GUTTER]
TOP, BOTTOM = 8.0, 34.5
BODY, SUBST, SMALL = 26, 24, 20  # body prompts; substantive boxes and tables; captions and references


def example_numbers() -> dict:
    """EDR|AI ch. 32's simulated checkout example (SEED = 464), with the contrast it is about."""
    import numpy as np
    from scipy import stats

    weeks = 12
    a = np.random.default_rng(SEED).normal(0.68, 0.070, size=weeks) * 100
    b = np.random.default_rng(SEED + 1).normal(0.705, 0.070, size=weeks) * 100
    ci = [1.96 * v.std(ddof=1) / np.sqrt(weeks) for v in (a, b)]  # as drawn in ch. 32
    va, vb = a.var(ddof=1) / weeks, b.var(ddof=1) / weeks
    se = np.sqrt(va + vb)
    df = (va + vb) ** 2 / (va**2 / (weeks - 1) + vb**2 / (weeks - 1))  # Welch, independent weeks
    t = stats.t.ppf(0.975, df)
    diff = b.mean() - a.mean()
    return {"weeks": weeks, "a": a.mean(), "b": b.mean(), "ci": ci,
            "diff": diff, "lo": diff - t * se, "hi": diff + t * se}


def num(x: float) -> str:
    """One decimal, with a true minus sign."""
    return f"{x:.1f}".replace("-", "\u2212")


def example_figure_png(n: dict) -> bytes:
    """The two panels a reader needs: each version's mean, and the difference the headline is about."""
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 22})
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(19, 11.0), dpi=200,
                                   gridspec_kw={"width_ratios": [1.15, 1]})
    means = [n["a"], n["b"]]
    ax1.bar(["Version A", "Version B"], means, yerr=n["ci"], capsize=14, width=0.55,
            color=["#4a4a4a", "#a8a8a8"], hatch=["", "//"], edgecolor="black",
            error_kw={"elinewidth": 2.5, "capthick": 2.5})
    ax1.set_ylim(0, 100)
    ax1.set_ylabel("Checkout completion rate (%)")
    ax1.set_title("(a) Each version: mean and 95% interval", fontsize=24, loc="left")
    for i, (m, e) in enumerate(zip(means, n["ci"])):
        ax1.text(i, m + e + 3, f"{m:.0f}% ± {e:.0f}", ha="center", fontsize=24, fontweight="bold")

    ax2.axvline(0, color="black", linewidth=2, linestyle="--")
    ax2.errorbar([n["diff"]], [0], xerr=[[n["diff"] - n["lo"]], [n["hi"] - n["diff"]]], fmt="o",
                 color="black", markersize=16, capsize=16, elinewidth=3, capthick=3)
    ax2.set_xlim(-10, 15)
    ax2.set_ylim(-1, 1)
    ax2.set_yticks([])
    ax2.set_xlabel("Difference, B minus A (percentage points)")
    ax2.set_title("(b) The difference: estimate and 95% interval", fontsize=24, loc="left")
    ax2.text(n["diff"] + 0.8, 0.3, f"{num(n['diff'])} points\n[{num(n['lo'])}, {num(n['hi'])}]",
             ha="left", fontsize=24, fontweight="bold")
    ax2.text(0.3, -0.75, "no difference", fontsize=20, color="#333333")
    for ax in (ax1, ax2):
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
    ax2.spines["left"].set_visible(False)
    fig.tight_layout(w_pad=3)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200)
    plt.close(fig)
    return buf.getvalue()


def build_template(qm474: Path, out: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.util import Inches, Pt

    # The Purdue mark from QM 47400's template; its Daniels School half is cropped.
    src = Presentation(str(qm474))
    logo = next(
        sh.image.blob
        for sh in src.slides[0].shapes
        if sh.shape_type == 13 and sh.image.size == (1423, 152)
    )
    ex = example_numbers()

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W_IN), Inches(H_IN)
    blank = prs.slide_layouts[6]
    rgb = RGBColor.from_string

    def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, name=None):
        s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if fill:
            s.fill.solid()
            s.fill.fore_color.rgb = rgb(fill)
        else:
            s.fill.background()
        if line:
            s.line.color.rgb = rgb(line)
            s.line.width = Pt(2.5)
        else:
            s.line.fill.background()
        s.shadow.inherit = False
        if name:
            s.name = name
        return s

    def text(slide, x, y, w, h, paras, name=None, anchor=MSO_ANCHOR.TOP, inset=0.12, shape=None, size=BODY):
        """paras: list of paragraphs; each a list of (text, opts) runs, or a str."""
        if shape is None:
            shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shape.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = anchor
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, Inches(inset))
        for k, para in enumerate(paras):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            if isinstance(para, str):
                para = [(para, {})]
            popts = para[0][1] if para else {}
            p.alignment = popts.get("align", PP_ALIGN.LEFT)
            p.space_after = Pt(popts.get("after", round(size * 0.4)))
            for t, o in para:
                r = p.add_run()
                r.text = t
                f = r.font
                f.name = FONT
                f.size = Pt(o.get("size", size))
                f.bold = o.get("bold", False)
                f.italic = o.get("italic", False)
                f.color.rgb = rgb(o.get("color", INK))
        if name:
            shape.name = name
        return shape

    def bar(slide, x, w, y, title, size=34):
        b = rect(slide, x, y, w, 0.95, INK, name=f"Heading: {title}")
        text(slide, 0, 0, 0, 0, [[(title, {"size": size, "bold": True, "color": GOLD, "after": 0})]],
             anchor=MSO_ANCHOR.MIDDLE, inset=0.3, shape=b)
        return y + 0.95

    def table(slide, x, y, w, rows, colw, size=SUBST, header=True, row_h=0.7, name=None):
        nrows, ncols = len(rows), len(rows[0])
        gt = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h * nrows))
        gt.name = name or "Table"
        tbl = gt.table
        tbl.first_row = False
        tbl.horz_banding = False
        for j, cw in enumerate(colw):
            tbl.columns[j].width = Inches(cw)
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = tbl.cell(i, j)
                cell.fill.solid()
                is_head = header and i == 0
                cell.fill.fore_color.rgb = rgb(INK if is_head else WHITE)
                cell.margin_left = cell.margin_right = Inches(0.12)
                cell.margin_top = cell.margin_bottom = Inches(0.06)
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.text = ""
                r = p.add_run()
                r.text = val
                r.font.name = FONT
                r.font.size = Pt(size)
                r.font.bold = is_head or (j == 0 and not header)
                r.font.color.rgb = rgb(GOLD if is_head else INK)
        return gt

    def labeled_box(slide, x, y, w, h, label, body, label_w=3.1, name=None):
        lb = rect(slide, x, y, label_w, h, INK, name=f"{name}: label")
        text(slide, 0, 0, 0, 0, [[(label, {"size": 22, "bold": True, "color": GOLD, "align": PP_ALIGN.CENTER,
                                           "after": 0})]], anchor=MSO_ANCHOR.MIDDLE, inset=0.1, shape=lb)
        bx = rect(slide, x + label_w, y, w - label_w, h, WHITE, line=INK, name=name)
        text(slide, 0, 0, 0, 0, [[(body, {"size": SUBST, "after": 0})]], anchor=MSO_ANCHOR.MIDDLE, inset=0.2,
             shape=bx)

    # ------------------------------------------------------------- slide 1: the poster
    s = prs.slides.add_slide(blank)
    left_edge = COL_X[0] + COL_W[0] + GUTTER / 2
    right_edge = COL_X[2] - GUTTER / 2
    rect(s, 0, TOP - 0.35, left_edge, H_IN - TOP + 0.35, GREY, name="Panel: left column")
    rect(s, right_edge, TOP - 0.35, W_IN - right_edge, H_IN - TOP + 0.35, GREY, name="Panel: right column")

    # header: title, authors, Purdue mark + Honors College, draft tag
    text(s, MARGIN, 0.35, 31.5, 2.55,
         [[("Title: your research problem and its answer", {"size": 72, "bold": True, "after": 0})]],
         name="Title")
    text(s, MARGIN, 2.95, 31.5, 1.5, [
        [("Your Name", {"size": 36, "bold": True}),
         ("  ·  Your major, Your college  ·  John Martinson Honors College, Purdue University", {"size": 32})],
        [("Faculty mentor: Professor Davi Moreira, Mitch Daniels School of Business  ·  your.email@purdue.edu",
          {"size": 28, "color": MUTED})],
    ], name="Authors and affiliations")
    logo_h = 1.75
    pic = s.shapes.add_picture(io.BytesIO(logo), Inches(33.2), Inches(0.75), height=Inches(logo_h))
    pic.crop_right = 1 - 778 / 1423
    pic.width = Inches(logo_h * 778 / 152)
    pic.name = "Purdue University mark"
    pic._element.nvPicPr.cNvPr.set("descr", "Purdue University")
    x_rule = 33.2 + logo_h * 778 / 152 + 0.35
    rect(s, x_rule, 0.8, 0.04, logo_h - 0.1, "8E8E8E", name="Lockup rule")
    text(s, x_rule + 0.25, 0.7, W_IN - MARGIN - x_rule - 0.25, 1.9,
         [[("John Martinson", {"size": 34, "bold": True, "after": 0})],
          [("Honors College", {"size": 34, "bold": True})]],
         anchor=MSO_ANCHOR.MIDDLE, inset=0.05, name="Honors College")
    tag = rect(s, 37.75, 3.0, W_IN - MARGIN - 37.75, 1.05, CRIMSON, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
               name="Draft tag (delete at the M13 lock)")
    text(s, 0, 0, 0, 0, [[("PRELIMINARY DRAFT · version 1", {"size": 30, "bold": True, "color": WHITE,
                                                              "align": PP_ALIGN.CENTER, "after": 0})]],
         anchor=MSO_ANCHOR.MIDDLE, shape=tag)

    # headline banner: the finding and the qualifier that changes how it reads, in one eye-span (ch. 32)
    hb = rect(s, MARGIN, 4.55, W_IN - 2 * MARGIN, 2.95, INK, name="Headline banner")
    text(s, 0, 0, 0, 0, [
        [("HEADLINE  ", {"size": 32, "bold": True, "color": GOLD}),
         ("Your main finding, with the uncertainty or boundary that changes how a reader should take it.",
          {"size": 46, "bold": True, "color": WHITE, "after": 6})],
        [(f"Simulated example: checkout version B averaged {num(ex['diff'])} points higher than version A, but "
          f"the 95% interval for that difference runs from {num(ex['lo'])} to {num(ex['hi'])} points "
          f"({ex['weeks']} weekly samples per version, one store).",
          {"size": 32, "italic": True, "color": GOLD})],
    ], anchor=MSO_ANCHOR.MIDDLE, inset=0.45, shape=hb)

    # ---- left column: the question; the design; the data
    x, w = COL_X[0], COL_W[0]
    y = bar(s, x, w, TOP, "THE QUESTION AND WHAT IT ADDS")
    text(s, x, y + 0.1, w, 8.3, [
        [("Your research question,", {"bold": True}), (" in one sentence that ends with a question mark, "
                                                       "and why it matters, and to whom.", {})],
        [("What the literature already establishes,", {"bold": True}),
         (" from at least two sources you read yourself (Author, Year).", {})],
        [("What your project adds", {"bold": True}), (" to that conclusion.", {})],
        [("Your main hypothesis,", {"bold": True}),
         (" as you recorded it before the analysis ran. Label any analysis you added later as exploratory.", {})],
    ], name="Question text")
    y = bar(s, x, w, 17.55, "DESIGN: WHAT YOUR CLAIM NEEDS")
    rows = [
        ("KIND", "☐ Descriptive   ☐ Causal.  If your aim is prediction, say so here."),
        ("REACH", "☐ A population beyond your data: name its units, setting, and period.  "
                  "☐ Unseen cases: which ones, from which population."),
        ("TARGET", "The exact quantity you estimate or predict: outcome, comparison, units, time frame."),
        ("WARRANT", "What carries the claim across: sampling → a population · assignment or identification "
                    "(with its assumptions) → a cause · a held-out test against a baseline → unseen cases. "
                    "Then: what your evidence cannot yet support."),
    ]
    yy = y + 0.2
    heights = [1.2, 1.75, 1.75, 2.75]
    for (label, body), h in zip(rows, heights):
        labeled_box(s, x, yy, w, h, label, body, name=f"Compass box: {label}")
        yy += h + 0.15
    y = bar(s, x, w, yy + 0.25, "DATA AND MEASUREMENT")
    card = [
        ("Source and route", "[archive, URL or DOI, retrieval date; or your collection protocol]"),
        ("Observations", "[number of rows or measurements]"),
        ("Independent units", "[people, sites, groups …] and how observations are grouped"),
        ("Time window", "[from … to …]"),
        ("Concept → measure", "[concept] → [variable or item]; main validity concern"),
        ("Missing data", "[how much, and what you did]"),
        ("Permissions", "[data licence; IRB determination]"),
    ]
    table(s, x, y + 0.2, w, card, [3.6, w - 3.6], size=SUBST, header=False, row_h=0.72, name="Data card")

    # ---- centre column: the main result, then the checks
    x, w = COL_X[1], COL_W[1]
    y = bar(s, x, w, TOP, "MAIN RESULT")
    text(s, x, y + 0.1, w, 0.65, [[("EXAMPLE · replace it with the figure your notebook produced, and "
                                     "replace its alt text with yours", {"size": 22, "bold": True,
                                                                         "color": CRIMSON, "after": 0})]],
         name="Example stamp")
    fig_h = w * 11.0 / 19
    fig = s.shapes.add_picture(io.BytesIO(example_figure_png(ex)), Inches(x), Inches(y + 0.8), width=Inches(w))
    fig.name = "Example figure (replace)"
    fig._element.nvPicPr.cNvPr.set(
        "descr",
        f"Simulated example. Left: mean checkout completion rate by version, A about {ex['a']:.0f} percent and "
        f"B about {ex['b']:.0f} percent, each with a 95 percent interval of about plus or minus 4 points. "
        f"Right: the difference, B minus A, {num(ex['diff'])} points, with a 95 percent interval from "
        f"{num(ex['lo'])} to {num(ex['hi'])} points that crosses zero.",
    )
    fy = y + 0.8 + fig_h
    text(s, x, fy + 0.1, w, 1.85, [
        [("Fig. 1.", {"size": SMALL, "bold": True}),
         (" Say what the figure shows and for which units; n = [observations] from [independent units]; the "
          "interval's type, level and method; bars start at zero. Source: [notebook], cell [N]. "
          "In the example, panel (b) carries the comparison the headline is about: the interval of the "
          "difference, not whether the two intervals in (a) overlap.", {"size": SMALL})],
    ], name="Figure caption")
    text(s, x, fy + 2.05, w, 2.0, [
        [("The result,", {"bold": True}),
         (" in one or two sentences: the estimate, its units, and the uncertainty it warrants (an interval, "
          "or a precise caveat saying why none is warranted), in the words your boundary allows.", {})],
    ], name="Result text")
    y = bar(s, x, w, 28.3, "CHECKS AND VERIFICATION")
    text(s, x, y + 0.1, w, BOTTOM - y - 0.1, [
        [("Analysis:", {"bold": True}),
         (" what you declared before seeing results and why it fits your design; any later change, and why.", {})],
        [("Uncertainty:", {"bold": True}),
         (" how you computed it, at the level of your independent unit, with any dependence named.", {})],
        [("One check:", {"bold": True}), (" a robustness, sensitivity, or diagnostic check, and what it changed.",
                                         {})],
    ], name="Checks text")

    # ---- right column: meaning and boundary; sources, AI use, credits
    x, w = COL_X[2], COL_W[2]
    y = bar(s, x, w, TOP, "WHAT IT MEANS, AND WHERE IT STOPS", size=32)
    for k, (label, body) in enumerate([
        ("THIS EVIDENCE SUPPORTS", "[your claim, with its verb and its population]"),
        ("IT DOES NOT YET SUPPORT", "[the larger claim a visitor might infer: a cause, a wider population, "
                                    "a longer time]"),
    ]):
        bx = rect(s, x, y + 0.2 + k * 2.45, w, 2.3, WHITE, line=INK, name=f"Boundary: {label}")
        text(s, 0, 0, 0, 0, [[(label, {"size": 22, "bold": True, "after": 4})],
                             [(body, {"size": SUBST, "after": 0})]], inset=0.2, shape=bx)
    text(s, x, y + 5.25, w, 21.75 - y - 5.25, [
        [("Answer:", {"bold": True}), (" your question, in one sentence, no more strongly than the evidence allows.",
                                      {})],
        [("Uncertainty:", {"bold": True}), (" how far the estimate could move.", {})],
        [("Limitation:", {"bold": True}), (" the one most likely to change the finding.", {})],
        [("Next step:", {"bold": True}), (" the study that would settle it.", {})],
    ], name="Meaning and ULN text")
    y = bar(s, x, w, 22.05, "SOURCES, AI USE, AND CREDITS", size=32)
    text(s, x, y + 0.1, w, 4.3, [
        [("AI use, on the poster itself.", {"bold": True, "size": SUBST}),
         (" One line per tool:", {"size": SUBST})],
        [("[Tool]: used to [purpose]; I [changed / kept what]; I checked it by [method].",
          {"size": SUBST, "italic": True})],
        [("Write “No AI used” only if that is true.", {"size": SMALL, "color": MUTED})],
    ], name="AI disclosure")
    text(s, x, y + 4.5, w, 3.3, [
        [("References.", {"bold": True, "size": SMALL}),
         (" Three to six sources you opened yourself: Author (Year). Title. Venue. DOI or URL.", {"size": SMALL})],
    ], name="References text")
    text(s, x, y + 7.9, w - 2.75, BOTTOM - y - 7.9, [
        [("Acknowledgments.", {"bold": True, "size": 18}),
         (" I thank Professor Davi Moreira for mentoring this project, and the John Martinson Honors College and "
          "the Office of Undergraduate Research for the chance to present it. [Thank any data provider, "
          "participants, or funder.]", {"size": 18})],
    ], name="Acknowledgments text")
    qr = rect(s, x + w - 2.55, y + 8.0, 2.55, 2.55, WHITE, line="8E8E8E", name="QR code placeholder")
    qr.line.dash_style = MSO_LINE.DASH
    text(s, 0, 0, 0, 0, [[("Optional QR + short URL: a text version of this poster",
                           {"size": 16, "color": MUTED, "align": PP_ALIGN.CENTER, "after": 0})]],
         anchor=MSO_ANCHOR.MIDDLE, shape=qr)

    # footer: verification status and preliminary results
    text(s, MARGIN, 34.85, 27.5, 0.85, [[
        ("Verification status: [what has been checked, by whom, and when; what is still pending] · "
         "Poster version [N], [date]", {"size": 18, "color": MUTED, "after": 0})]], name="Verification status line")
    text(s, 28.5, 34.85, W_IN - MARGIN - 28.5, 0.85, [[
        ("Preliminary results · HONR 46400 Evidence-Driven Research · Purdue Fall Undergraduate Research Expo 2026",
         {"size": 18, "color": MUTED, "align": PP_ALIGN.RIGHT, "after": 0})]], name="Status line")

    # ------------------------------------------------------------- slide 2: hidden guide page
    g = prs.slides.add_slide(blank)
    g._element.set("show", "0")
    text(g, MARGIN, 0.6, W_IN - 2 * MARGIN, 2.2, [
        [("READ ME: how to use this template", {"size": 72, "bold": True})],
        [("This page is hidden, but some export settings still include hidden slides. "
          "Export slide 1 only, or delete this page first.", {"size": 34, "color": CRIMSON, "bold": True})],
    ], name="Guide title")
    gw = (W_IN - 2 * MARGIN - 2 * 1.0) / 3
    gx = [MARGIN + k * (gw + 1.0) for k in range(3)]

    def gbar(k, y, title):
        return bar(g, gx[k], gw, y, title)

    y = gbar(0, 4.2, "PRINT SPECIFICATION")
    text(g, gx[0], y + 0.1, gw, 18.0, [
        [("Size:", {"bold": True}), (" 48 × 36 in, landscape, the Expo's standard 4′ × 3′ board. "
                                      "Do not change the slide size.", {})],
        [("Type at print size:", {"bold": True}),
         (" title 72 pt or more; headline 44 pt; headings 32 pt; everything substantive, including the "
          "compass, the data card, the boundary boxes and the AI disclosure, 24 pt; captions, references, "
          "acknowledgments and footer 18 pt. If it does not fit, cut or reorganize. Never shrink below these.", {})],
        [("Figures:", {"bold": True}),
         (" save each one from your notebook as PNG at 300 dpi at the size it will print, or as PDF or SVG. "
          "Never paste a screenshot. Right-click → Edit Alt Text, and write what the figure shows.", {})],
        [("Export:", {"bold": True}),
         (" export slide 1 only, as a PDF. Open it and check: exactly one page, 48 × 36 in, nothing clipped, "
          "no missing figure. Then print one crop at actual size, without “fit to page”, and measure it with "
          "a ruler. Screen zoom does not tell you the printed size.", {})],
        [("Files:", {"bold": True}),
         (" M10: lastname_m10_poster.pptx and lastname_m10_poster.pdf. M11: lastname_m11_posterdraft.pdf. "
          "M13: lastname_m13_poster.pdf, the print file, poster only.", {})],
    ], name="Guide: print spec", size=34)

    y = gbar(1, 4.2, "RULES THE BOARD KEEPS")
    text(g, gx[1], y + 0.1, gw, 29.5, [
        [("No result appears first on the poster.", {"bold": True}),
         (" Every result, figure and citation comes from your M9 note and its claim-evidence table. If a check "
          "changes a number, keep the submitted note, write a dated corrected version, then update the poster.",
          {})],
        [("Trace what you print.", {"bold": True}),
         (" A result you computed traces to its notebook cell; a statistic from a source traces to that source; "
          "rounding is documented.", {})],
        [("The qualifier travels with the claim,", {"bold": True}),
         (" in the same eye-span as the headline.", {})],
        [("Causal wording only with a defended identification argument;", {"bold": True}),
         (" population wording only with a reach argument.", {})],
        [("Two channels for every distinction:", {"bold": True}), (" a label or a pattern, with colour as a bonus.",
                                                                    {})],
        [("Replace every bracket, every prompt and every EXAMPLE item.", {"bold": True}),
         (" Panels may move, merge or resize; every kind of information stays.", {})],
        [("Draft tag through M11; delete it at the M13 lock.", {"bold": True}),
         (" Keep “Preliminary results” in the footer: the project closes in December.", {})],
        [("AI:", {"bold": True}),
         (" disclose on the poster itself, per tool: purpose, what you changed or kept, how you checked it. The "
          "Expo does not accept AI-generated sections or abstracts, or a poster designed by AI without "
          "substantial human input. Every AI step also gets a ledger row.", {})],
    ], name="Guide: rules", size=34)

    y = gbar(2, 4.2, "PANEL → RUBRIC → BOOK")
    table(g, gx[2], y + 0.2, gw, [
        ("Panel", "Rubric", "Book"),
        ("Title · The question and what it adds", "1", "Ch. 1, 7–9"),
        ("Design: what your claim needs", "2", "Ch. 6–7, 10, 12, 14–19"),
        ("Data and measurement", "3", "Ch. 13, 20–21"),
        ("Checks and verification · status line", "4", "Ch. 11, 22–25, 36–37"),
        ("Main result", "5", "Ch. 32"),
        ("Headline · What it means, and where it stops", "6", "Ch. 28–29, 35"),
        ("Sources, AI use, and credits", "7", "Ch. 9, 30"),
        ("The whole board", "8", "Ch. 32–33"),
    ], [gw - 6.4, 2.0, 4.4], size=28, row_h=1.1, name="Guide: panel map")
    text(g, gx[2], 17.5, gw, 8.0, [
        "The final poster (M13) is graded with Evidence_Driven_Research_Poster_Rubric.pdf, handed out with "
        "this template: eight criteria, 100 points, with hard caps. Read it before you build, not after.",
    ], name="Guide: rubric", size=34)

    prs.core_properties.title = "HONR 46400 research poster template"
    prs.core_properties.author = "Davi Moreira"
    prs.core_properties.subject = "Purdue Fall Undergraduate Research Expo poster, 48 x 36 in"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--rubric", action="store_true", help="rubric documents and the M13 block only")
    ap.add_argument("--check", action="store_true", help="exit 1 if the M13 rubric block is stale")
    ap.add_argument("--qm474", type=Path, default=QM474_DEFAULT, help="QM 47400 template (logo source)")
    args = ap.parse_args()

    if args.check:
        ok = inject_m13(check=True)
        print("M13 poster rubric block:", "current" if ok else "STALE (run build_poster_kit.py)")
        return 0 if ok else 1

    inject_m13()
    docx_path = OUT_DIR / f"{RUBRIC_STEM}.docx"
    build_rubric_docx(docx_path)
    to_pdf(docx_path, OUT_DIR)
    build_rubric_xlsx(OUT_DIR / f"{RUBRIC_STEM}.xlsx")
    print(f"wrote {docx_path.relative_to(ROOT)} (+ .pdf, .xlsx); M13 block current")
    if not args.rubric:
        if not args.qm474.exists():
            sys.exit(f"QM 47400 template not found: {args.qm474} (pass --qm474 PATH)")
        build_template(args.qm474, TEMPLATE_OUT)
        print(f"wrote {TEMPLATE_OUT.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
